"""Export Data tab: teacher/admin field-picker export to Excel/PDF.

Covers the /export page + POST /api/export/students route:
  * catalog      — form.html-pane grouping, core/record groups, no duplicate
                   columns, labels present, Yes/No flags, defaults in catalog
  * page guards  — anon 302, student 403, teacher/admin 200 with picker
  * api guards   — anon 401, student 403, format/fields validation -> 422
  * scoping      — teacher xlsx = own class+section rows only; admin = all
  * content      — xlsx headers = requested labels, Yes/No rendering,
                   pdf = %PDF magic, download filenames carry the scope tag
  * invariants   — export is read-only (student ids untouched), cell_value
                   unit rules (bool/1/0/datetime/None)

Fixtures are created (2 students) and removed again in finally; the
baseline id list is captured BEFORE fixture creation so the check works
against any restored DB state.

Run: python test_export.py   (exit 0 = all pass)
"""
import io
import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT / "femis-web"))
sys.path.insert(0, str(ROOT))

from app import (  # noqa: E402
    app, db, Student,
    _EXPORT_GROUPS, _EXPORT_LABELS, _EXPORT_YESNO, _EXPORT_DEFAULTS,
)
import student_export  # noqa: E402
from form_payloads import tab1  # noqa: E402
from sqlalchemy import text  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DB = ROOT / "femis-web" / "instance" / "femis.db"
TEACHER_CLASS, TEACHER_SECTION = "9", "A"
OTHER_CLASS, OTHER_SECTION = "8", "B"

results = []


def check(name, ok, detail=""):
    results.append((name, ok, detail))
    print(f"{'PASS' if ok else 'FAIL'}: {name}" + (f" — {detail}" if detail else ""))


def client(role=None, name="exporttest"):
    c = app.test_client()
    if role:
        with c.session_transaction() as s:
            s["role"] = role
            s["user_name"] = name
            if role == "teacher":
                s["teacher_class"] = TEACHER_CLASS
                s["teacher_section"] = TEACHER_SECTION
    return c


def create_student(name, class_id, section_id, roll):
    r = app.test_client().post(
        "/api/save-tab",
        json={"tab": 1, "student_id": None,
              "data": tab1(name=name, class_id=class_id,
                           section_id=section_id, roll_no=roll)},
    )
    b = r.get_json() or {}
    return b.get("student_id") if b.get("ok") else None


def sheet_rows(resp):
    wb = load_workbook(io.BytesIO(resp.data))
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


created_students = []
baseline_ids = []

try:
    with app.app_context():
        baseline_ids = [i for (i,) in db.session.execute(
            text("SELECT id FROM students ORDER BY id"))]

    # ------------------------------------------------------------------
    # A. Catalog structure
    # ------------------------------------------------------------------
    titles = [t for t, _ in _EXPORT_GROUPS]
    all_cols = [c for _, items in _EXPORT_GROUPS for c, _ in items]
    stu_cols = {c.name for c in Student.__table__.columns}
    check("A1 groups: Student first, Record last, >= 8 groups",
          len(titles) >= 8 and titles[0] == "Student" and titles[-1] == "Record",
          titles)
    check("A2 no duplicate columns; role excluded; covers all real columns",
          len(all_cols) == len(set(all_cols))
          and "role" not in all_cols
          and set(all_cols) == stu_cols - {"role"},
          f"{len(all_cols)} cols vs {len(stu_cols)} columns")
    check("A3 every column has a non-empty label",
          all(isinstance(l, str) and l.strip() for l in _EXPORT_LABELS.values())
          and len(_EXPORT_LABELS) == len(all_cols),
          len(_EXPORT_LABELS))
    check("A4 defaults = 11 core fields, all present in catalog",
          _EXPORT_DEFAULTS == [c for c, _ in student_export._CORE_FIELDS]
          and all(c in _EXPORT_LABELS for c in _EXPORT_DEFAULTS),
          _EXPORT_DEFAULTS)
    gmap = {t: [c for c, _ in items] for t, items in _EXPORT_GROUPS}
    check("A5 pane grouping: emergency_* under Emergency, digital_* under Digital",
          "emergency_contact" in gmap.get("Emergency Contact", [])
          and "emergency_name" in gmap.get("Emergency Contact", [])
          and "digital_device_type" in gmap.get("Digital Access", [])
          and "father_cnic" in gmap.get("Parents / Guardian", []),
          {t: len(v) for t, v in gmap.items()})
    check("A6 Yes/No flags: submitted/locked + form yesno cols, all in catalog",
          {"submitted", "locked"} <= _EXPORT_YESNO
          and "digital_device_at_home" in _EXPORT_YESNO
          and _EXPORT_YESNO <= set(_EXPORT_LABELS),
          sorted(_EXPORT_YESNO)[:6])
    check("A7 catalog covers form panes (7 pane groups + core + record)",
          len(titles) == 9, titles)

    # ------------------------------------------------------------------
    # B. GET /export page guards
    # ------------------------------------------------------------------
    anon = app.test_client()
    r = anon.get("/export", follow_redirects=False)
    check("B1 anon /export -> login redirect",
          r.status_code == 302 and "/login" in (r.headers.get("Location") or ""),
          r.headers.get("Location"))

    stu = client(role="student")
    r = stu.get("/export")
    check("B2 student /export -> 403", r.status_code == 403, r.status_code)

    tch = client(role="teacher")
    r = tch.get("/export")
    body = r.get_data(as_text=True)
    scope_txt = f"Class {TEACHER_CLASS} - Section {TEACHER_SECTION}"
    check("B3 teacher /export -> 200 with scope badge + field picker",
          r.status_code == 200 and scope_txt in body
          and 'id="exportFields"' in body and "Export Excel" in body
          and "Export PDF" in body,
          str(r.status_code))

    adm = client(role="admin")
    r = adm.get("/export")
    body = r.get_data(as_text=True)
    check("B4 admin /export -> 200, All classes scope",
          r.status_code == 200 and "All classes" in body
          and 'id="exportFields"' in body,
          str(r.status_code))

    # ------------------------------------------------------------------
    # C. POST /api/export/students guards + validation
    # ------------------------------------------------------------------
    r = anon.post("/api/export/students",
                  json={"format": "xlsx", "fields": ["name"]})
    b = r.get_json() or {}
    check("C1 anon api -> 401 authentication_required",
          r.status_code == 401 and not b.get("ok"), r.status_code)

    r = stu.post("/api/export/students",
                 json={"format": "xlsx", "fields": ["name"]})
    check("C2 student api -> 403", r.status_code == 403, r.status_code)

    for label, payload in [
        ("bad format", {"format": "csv", "fields": ["name"]}),
        ("missing fields", {"format": "xlsx"}),
        ("empty fields list", {"format": "xlsx", "fields": []}),
        ("non-list fields", {"format": "xlsx", "fields": "name"}),
        ("only non-string entries", {"format": "xlsx", "fields": [1, None]}),
    ]:
        r = adm.post("/api/export/students", json=payload)
        b = r.get_json() or {}
        check(f"C3 validation: {label} -> 422",
              r.status_code == 422 and not b.get("ok") and b.get("error"),
              f"{r.status_code} {b.get('error')}")

    r = adm.post("/api/export/students",
                 json={"format": "xlsx", "fields": ["name", "not_a_column"]})
    b = r.get_json() or {}
    check("C4 unknown field named in 422 error",
          r.status_code == 422 and "not_a_column" in (b.get("error") or ""),
          b.get("error"))

    too_many = list(_EXPORT_LABELS)[:31]
    r = adm.post("/api/export/students",
                 json={"format": "pdf", "fields": too_many})
    b = r.get_json() or {}
    check("C5 pdf >30 columns -> 422 suggesting Excel",
          r.status_code == 422 and "Excel" in (b.get("error") or ""),
          b.get("error"))

    # ------------------------------------------------------------------
    # D. Fixtures (own class + foreign class) for scoping checks
    # ------------------------------------------------------------------
    sid_own = create_student("Export Own Student", TEACHER_CLASS,
                             TEACHER_SECTION, "9091")
    sid_other = create_student("Export Other Student", OTHER_CLASS,
                               OTHER_SECTION, "9092")
    created_students += [sid_own, sid_other]
    check("D1 fixtures created", bool(sid_own and sid_other),
          (sid_own, sid_other))

    fields = ["name", "class_id", "section_id", "roll_no", "submitted", "locked"]
    labels = [_EXPORT_LABELS[c] for c in fields]

    # ------------------------------------------------------------------
    # E. Teacher export: scoped rows only, correct headers/filename
    # ------------------------------------------------------------------
    with app.app_context():
        own_names = [n for (n,) in db.session.execute(text(
            "SELECT name FROM students WHERE class_id = :c AND section_id = :s"),
            {"c": TEACHER_CLASS, "s": TEACHER_SECTION})]
        total = db.session.execute(text(
            "SELECT COUNT(*) FROM students")).scalar()

    r = tch.post("/api/export/students",
                 json={"format": "xlsx", "fields": fields})
    cd = r.headers.get("Content-Disposition") or ""
    rows = sheet_rows(r) if r.status_code == 200 else []
    check("E1 teacher xlsx 200 + filename students_9A_<date>.xlsx",
          r.status_code == 200 and "students_9A_" in cd and ".xlsx" in cd,
          f"{r.status_code} {cd}")
    check("E2 xlsx headers == requested labels",
          rows and list(rows[0]) == labels,
          rows[0] if rows else None)
    body_names = {row[0] for row in rows[1:]}
    check("E3 teacher xlsx = own class rows only (foreign student absent)",
          len(rows) - 1 == len(own_names)
          and "Export Other Student" not in body_names
          and "Export Own Student" in body_names,
          f"rows={len(rows) - 1} expected={len(own_names)}")

    # ------------------------------------------------------------------
    # F. Admin export: all rows
    # ------------------------------------------------------------------
    r = adm.post("/api/export/students",
                 json={"format": "xlsx", "fields": fields})
    cd = r.headers.get("Content-Disposition") or ""
    rows = sheet_rows(r) if r.status_code == 200 else []
    body_names = {row[0] for row in rows[1:]}
    check("F1 admin xlsx 200 + filename students_all_<date>.xlsx",
          r.status_code == 200 and "students_all_" in cd and ".xlsx" in cd,
          f"{r.status_code} {cd}")
    check("F2 admin xlsx covers every student (both fixtures present)",
          len(rows) - 1 == total
          and {"Export Own Student", "Export Other Student"} <= body_names,
          f"rows={len(rows) - 1} total={total}")

    # ------------------------------------------------------------------
    # G. Yes/No rendering + PDF
    # ------------------------------------------------------------------
    sub_row = [row for row in rows[1:] if row[0] == "Export Own Student"]
    yesno_vals_ok = (sub_row and sub_row[0][4] in ("Yes", "No")
                     and sub_row[0][5] in ("Yes", "No"))
    check("G1 submitted/locked render as Yes/No", bool(yesno_vals_ok),
          sub_row[0][4:6] if sub_row else None)

    r = adm.post("/api/export/students",
                 json={"format": "pdf", "fields": ["name", "class_id",
                                                   "section_id", "roll_no"]})
    cd = r.headers.get("Content-Disposition") or ""
    check("G2 admin pdf 200, %PDF magic, .pdf filename",
          r.status_code == 200 and r.data[:4] == b"%PDF"
          and ".pdf" in cd and "students_all_" in cd,
          f"{r.status_code} {r.data[:4]!r} {cd}")

    r = tch.post("/api/export/students",
                 json={"format": "pdf", "fields": ["name", "roll_no"]})
    check("G3 teacher pdf 200 (scoped)",
          r.status_code == 200 and r.data[:4] == b"%PDF"
          and "students_9A_" in (r.headers.get("Content-Disposition") or ""),
          r.status_code)

    # G4: a pdf of 100 rows must span multiple pages — with auto page break
    # off the table rendered ONE page and silently dropped everything past
    # row ~22 (live bug: excel had all students, pdf had one page).
    fake_rows = [SimpleNamespace(name=f"Pagefill Student {i:03d}",
                                 class_id="7", section_id="A", roll_no=str(i))
                 for i in range(1, 101)]
    buf = student_export.students_to_pdf(
        fake_rows, ["name", "class_id", "section_id", "roll_no"],
        ["Name", "Class", "Section", "Roll No"], "Student List - Test")
    pages = len(re.findall(rb"/Type\s*/Page(?!s)", buf.getvalue()))
    check("G4 pdf paginates: 100 rows span >= 2 pages (no dropped rows)",
          pages >= 2, f"{pages} page(s)")

    # ------------------------------------------------------------------
    # H. cell_value unit rules
    # ------------------------------------------------------------------
    fake = SimpleNamespace(submitted=1, locked=0, name="X",
                           created_at=datetime(2026, 10, 5, 16, 30),
                           date_of_birth=date(2026, 1, 2),
                           father_name=None)
    check("H1 bool/1/0 -> Yes/No (yesno col), raw otherwise",
          student_export.cell_value(fake, "submitted", True) == "Yes"
          and student_export.cell_value(fake, "locked", True) == "No"
          and student_export.cell_value(fake, "submitted", False) == "1",
          (student_export.cell_value(fake, "submitted", True),
           student_export.cell_value(fake, "locked", True),
           student_export.cell_value(fake, "submitted", False)))
    check("H2 datetime/date formatted, None -> empty string",
          student_export.cell_value(fake, "created_at") == "2026-10-05 16:30"
          and student_export.cell_value(fake, "date_of_birth") == "2026-01-02"
          and student_export.cell_value(fake, "father_name") == "",
          (student_export.cell_value(fake, "created_at"),
           student_export.cell_value(fake, "date_of_birth")))

    # ------------------------------------------------------------------
    # I. Dashboard buttons (entry points)
    # ------------------------------------------------------------------
    dash = tch.get("/teacher-dashboard").get_data(as_text=True)
    check("I1 teacher dashboard links /export",
          'href="/export"' in dash, dash.count('href="/export"'))
    adash = adm.get("/teacher-dashboard").get_data(as_text=True)
    check("I2 admin dashboard links /export",
          'href="/export"' in adash, adash.count('href="/export"'))

    check("I3 export page has WHERE controls + Run Query + result table",
          'id="filterClass"' in body and 'id="runQuery"' in body
          and 'id="resultTable"' in body and 'id="perPage"' in body
          and 'id="btnXlsx"' in body,
          "marker scan")

    # ------------------------------------------------------------------
    # J. Preview API: guards + payload validation
    # ------------------------------------------------------------------
    r = anon.post("/api/export/preview", json={"fields": ["name"]})
    check("J1 anon preview -> 401", r.status_code == 401, r.status_code)
    r = stu.post("/api/export/preview", json={"fields": ["name"]})
    check("J2 student preview -> 403", r.status_code == 403, r.status_code)

    r = tch.post("/api/export/preview", json={"fields": ["name"]})
    d = r.get_json() or {}
    check("J3 teacher preview 200 with labels/rows/scope",
          r.status_code == 200 and d.get("ok")
          and d.get("labels") == ["Name"]
          and isinstance(d.get("rows"), list)
          and d.get("scope") == f"Class {TEACHER_CLASS} - Section {TEACHER_SECTION}",
          (r.status_code, d.get("scope")))

    for label, payload, needle in [
        ("unknown filter key", {"fields": ["name"], "filters": {"bogus": "x"}},
         "bogus"),
        ("filters not an object", {"fields": ["name"], "filters": "class=7"},
         "object"),
        ("bad status value", {"fields": ["name"], "filters": {"status": "maybe"}},
         "status"),
        ("bad lock value", {"fields": ["name"], "filters": {"lock": "yes"}},
         "lock"),
        ("q too long", {"fields": ["name"], "filters": {"q": "x" * 101}},
         "100"),
        ("unknown field", {"fields": ["nope"]}, "nope"),
        ("bad sort", {"fields": ["name"], "sort": "bogus"}, "sort"),
        ("no fields", {"fields": []}, "at least one"),
    ]:
        r = adm.post("/api/export/preview", json=payload)
        b = r.get_json() or {}
        check(f"J4 preview validation: {label} -> 422",
              r.status_code == 422
              and needle.lower() in (b.get("error") or "").lower(),
              b.get("error"))

    r = adm.post("/api/export/preview", json={"fields": ["name"], "per": "all"})
    d = r.get_json() or {}
    check("J5 invalid per defaults to 25",
          r.status_code == 200 and d.get("per") == 25, d.get("per"))

    # ------------------------------------------------------------------
    # K. WHERE filters actually filter (+ teacher scope ceiling)
    # ------------------------------------------------------------------
    def db_col(sql, params):
        with app.app_context():
            return [r[0] for r in db.session.execute(text(sql), params)]

    r = adm.post("/api/export/preview",
                 json={"fields": ["name"], "filters": {"class": OTHER_CLASS}})
    d = r.get_json() or {}
    names = [row[0] for row in d.get("rows", [])]
    expected_other = db_col(
        "SELECT name FROM students WHERE class_id = :c", {"c": OTHER_CLASS})
    check("K1 class filter total == DB set; own excluded, other present",
          r.status_code == 200 and d.get("total") == len(expected_other)
          and "Export Other Student" in names
          and "Export Own Student" not in names,
          (d.get("total"), len(expected_other)))

    r = adm.post("/api/export/preview",
                 json={"fields": ["name"], "filters": {"q": "Export Own"}})
    d = r.get_json() or {}
    check("K2 text search narrows to the matching name",
          r.status_code == 200 and d.get("total") == 1
          and d.get("rows") == [["Export Own Student"]],
          (d.get("total"), d.get("rows")))

    r = adm.post("/api/export/preview",
                 json={"fields": ["name"],
                       "filters": {"q": "Export", "status": "draft"}})
    d = r.get_json() or {}
    draft_names = [row[0] for row in d.get("rows", [])]
    check("K3 status=draft contains both fresh fixtures",
          r.status_code == 200
          and {"Export Own Student", "Export Other Student"} <= set(draft_names),
          draft_names)

    r = adm.post("/api/export/preview",
                 json={"fields": ["name"],
                       "filters": {"q": "Export", "status": "submitted"}})
    d = r.get_json() or {}
    check("K4 status=submitted excludes the draft fixtures",
          r.status_code == 200 and d.get("total") == 0
          and d.get("rows") == [],
          d.get("total"))

    with app.app_context():
        db.session.execute(
            text("UPDATE students SET locked = 1 WHERE id = :i"),
            {"i": sid_own})
        db.session.commit()
    r_lock = adm.post("/api/export/preview",
                      json={"fields": ["name"],
                            "filters": {"q": "Export", "lock": "locked"}})
    r_unl = adm.post("/api/export/preview",
                     json={"fields": ["name"],
                           "filters": {"q": "Export", "lock": "unlocked"}})
    lock_names = [row[0] for row in (r_lock.get_json() or {}).get("rows", [])]
    unl_names = [row[0] for row in (r_unl.get_json() or {}).get("rows", [])]
    check("K5 lock filter separates locked vs unlocked fixtures",
          lock_names == ["Export Own Student"]
          and unl_names == ["Export Other Student"],
          (lock_names, unl_names))

    r = tch.post("/api/export/preview",
                 json={"fields": ["name"], "filters": {"class": OTHER_CLASS}})
    d = r.get_json() or {}
    check("K6 teacher foreign-class filter -> 0 rows (scope ceiling)",
          r.status_code == 200 and d.get("total") == 0
          and d.get("rows") == [],
          d.get("total"))

    expected_own = db_col(
        "SELECT name FROM students WHERE class_id = :c AND section_id = :s",
        {"c": TEACHER_CLASS, "s": TEACHER_SECTION})
    r = tch.post("/api/export/preview", json={"fields": ["name"]})
    d = r.get_json() or {}
    check("K7 teacher default preview total == own class+section count",
          r.status_code == 200 and d.get("total") == len(expected_own),
          (d.get("total"), len(expected_own)))

    # ------------------------------------------------------------------
    # L. Pagination + sort
    # ------------------------------------------------------------------
    r = adm.post("/api/export/preview",
                 json={"fields": ["name"], "per": 25, "page": 1})
    d = r.get_json() or {}
    total = d.get("total") or 0
    exp_pages = max(1, -(-total // 25))
    check("L1 page-1 math: pages, row count, per echoed",
          r.status_code == 200 and d.get("pages") == exp_pages
          and d.get("page") == 1 and d.get("per") == 25
          and len(d.get("rows") or []) == min(25, total),
          (total, d.get("pages"), len(d.get("rows") or [])))

    r = adm.post("/api/export/preview",
                 json={"fields": ["name"], "per": 25, "page": 999})
    d = r.get_json() or {}
    last_page_rows = total - (exp_pages - 1) * 25 if total else 0
    check("L2 out-of-range page clamps to last page",
          r.status_code == 200 and d.get("page") == exp_pages
          and len(d.get("rows") or []) == last_page_rows,
          (d.get("page"), exp_pages, len(d.get("rows") or [])))

    r = adm.post("/api/export/preview",
                 json={"fields": ["name"], "sort": "name"})
    d = r.get_json() or {}
    got = [row[0] for row in d.get("rows") or []]
    check("L3 sort=name returns rows A->Z (case-insensitive)",
          r.status_code == 200 and len(got) <= 25
          and got == sorted(got, key=lambda s: (s or "").casefold()),
          got[:5])

    # ------------------------------------------------------------------
    # M. Filtered downloads + preview/file parity (WYSIWYG)
    # ------------------------------------------------------------------
    r = adm.post("/api/export/students",
                 json={"format": "xlsx", "fields": ["name"],
                       "filters": {"class": OTHER_CLASS}})
    xnames = {row[0] for row in sheet_rows(r)[1:]}
    check("M1 filtered admin export == the filtered DB set",
          r.status_code == 200 and xnames == set(expected_other)
          and "Export Own Student" not in xnames,
          len(xnames))

    r = tch.post("/api/export/students",
                 json={"format": "xlsx", "fields": ["name"],
                       "filters": {"class": OTHER_CLASS}})
    mrows = sheet_rows(r) if r.status_code == 200 else [["bad"]]
    check("M2 teacher export + foreign filter -> header only (0 rows)",
          r.status_code == 200 and len(mrows) == 1,
          len(mrows))

    wfields = ["name", "class_id", "section_id", "roll_no"]
    p = adm.post("/api/export/preview",
                 json={"fields": wfields, "filters": {"q": "Export"},
                       "per": 100}).get_json() or {}
    x = adm.post("/api/export/students",
                 json={"format": "xlsx", "fields": wfields,
                       "filters": {"q": "Export"}})
    xrows = sheet_rows(x)
    check("M3 WYSIWYG: preview rows == downloaded xlsx rows",
          list(xrows[0]) == p.get("labels")
          and [[str(v) for v in row] for row in xrows[1:]] == p.get("rows"),
          (len(xrows) - 1, len(p.get("rows") or [])))

finally:
    # ------------------------------------------------------------------
    # Cleanup: temp students removed; baseline ids restored; DB read-only
    # (export never wrote a row — verified by the id compare).
    # ------------------------------------------------------------------
    with app.app_context():
        for sid in [s for s in created_students if s]:
            st = Student.query.get(sid)
            if st:
                db.session.delete(st)
        db.session.commit()

    con = sqlite3.connect(str(DB))
    ids = [r[0] for r in con.execute("SELECT id FROM students ORDER BY id")]
    con.close()
    check("Z1 students baseline ids restored after export suite",
          ids == baseline_ids, ids)
    check("Z2 both fixtures removed",
          all(s not in ids for s in created_students if s),
          (len(ids), len(baseline_ids)))

passed = sum(1 for _, ok, _ in results if ok)
print("=" * 50)
print(f"{passed}/{len(results)} checks passed")
sys.exit(0 if passed == len(results) else 1)
